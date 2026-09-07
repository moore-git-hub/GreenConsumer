from copy import deepcopy
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile
import shutil
import tempfile

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from lxml import etree


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "07-02 正文-李子业_第一轮修改_审阅版.docx"
OUTPUT = ROOT / "07-02 正文-李子业_第二轮机制统一版.docx"


def paragraph_after(paragraph: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    result = Paragraph(new_p, paragraph._parent)
    result.style = style
    result.add_run(text)
    return result


def find_heading(doc: Document, text: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    raise KeyError(f"Heading not found: {text}")


def replace_section(doc: Document, heading: str, next_heading: str, paragraphs: list[str]) -> None:
    start = find_heading(doc, heading)
    end = find_heading(doc, next_heading)
    node = start._p.getnext()
    while node is not None and node is not end._p:
        nxt = node.getnext()
        node.getparent().remove(node)
        node = nxt
    cursor = start
    for text in paragraphs:
        cursor = paragraph_after(cursor, text)


def replace_by_prefix(doc: Document, prefix: str, text: str) -> None:
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefix):
            p.text = text
            return
    raise KeyError(f"Paragraph not found: {prefix}")


def strip_comments(path: Path) -> None:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with tempfile.TemporaryDirectory() as td:
        temp_path = Path(td) / path.name
        with ZipFile(path, "r") as zin, ZipFile(temp_path, "w", ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                if info.filename.startswith("word/comments") and info.filename.endswith(".xml"):
                    continue
                data = zin.read(info.filename)
                if info.filename == "word/document.xml":
                    root = etree.fromstring(data)
                    for tag in ("commentRangeStart", "commentRangeEnd", "commentReference"):
                        for node in root.xpath(f"//w:{tag}", namespaces=ns):
                            node.getparent().remove(node)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
                elif info.filename == "word/_rels/document.xml.rels":
                    root = etree.fromstring(data)
                    for node in list(root):
                        if node.get("Type", "").endswith("/comments") or node.get("Target", "").startswith("comments"):
                            root.remove(node)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
                elif info.filename == "[Content_Types].xml":
                    root = etree.fromstring(data)
                    for node in list(root):
                        if node.get("PartName", "").startswith("/word/comments"):
                            root.remove(node)
                    data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone="yes")
                zout.writestr(info, data)
        shutil.copy2(temp_path, path)


def main() -> None:
    shutil.copy2(SOURCE, OUTPUT)
    doc = Document(OUTPUT)

    replace_section(doc, "消费者异质性类型", "状态变量与参数体系", [
        "消费者异质性是本研究构建智能体群体的基础设定。为使类型划分具有可核验的经验依据，本研究参考 Husson 等（2025）发布的《Forrester’s Green Consumer Segmentation For 2026》。该报告基于 2025 年消费者基准调查，分别考察美国、澳大利亚和欧洲五国成年在线消费者，并以三个问题识别消费者差异：是否主动搜寻绿色产品或绿色行为信息，是否愿意在较不便利的条件下仍购买绿色产品，以及环境因素是否让位于价格和便利性。",
        "依据上述分类逻辑，模型设置积极绿色消费者（Active Greens）、便利绿色消费者（Convenient Greens）、沉睡绿色消费者（Dormant Greens）和非绿色消费者（Non-Greens）四类 Agent。报告给出的区域占比分别为：欧洲五国 18%、27%、38% 和 17%；美国 13%、31%、35% 和 21%；澳大利亚 14%、27%、35% 和 24%。由于三个地区的构成并不相同，模型不能把任一组比例无条件解释为普遍人口分布。当前基准实验采用四类等额抽样，以提高类型间比较的统计可辨识度；若后续开展特定市场情景实验，再按相应地区比例进行加权抽样。",
        "积极绿色消费者会主动关注绿色产品及更环保的行为方式，并愿意在绿色选择较不便利时仍付诸行动。该群体对绿色产品的支付溢价意愿也相对较高。因此，模型以较高绿色信息关注和较强绿色选择倾向约束其 Persona。",
        "便利绿色消费者对环境议题具有一定兴趣，但在实际选择中更重视价格和便利性。模型据此将其设置为条件性绿色选择者：当绿色产品容易获得且价格可接受时更可能选择绿色方案，反之则更可能让位于现实成本。",
        "沉睡绿色消费者通常不主动搜寻绿色产品或绿色行为信息，也不会稳定地因环境属性选择产品，但在相关信息被明确呈现后可能受到说服。模型据此将其设置为低主动关注、可被信息激活的群体，但不预设其必然发生突变式态度转向。",
        "非绿色消费者在购买判断中将环境因素置于价格和便利性之后。模型据此降低绿色属性在其 Persona 中的决策权重，但不将其描述为对环境‘零关注’，也不预设其拒绝所有绿色溢价。",
        "Forrester 报告支持的是类型名称、分类问题、区域构成及一般绿色选择倾向，并未直接估计漂绿敏感度、初始品牌信任、信任恢复速率或发帖概率。因此，上述变量在模型中属于研究者设定的机制参数，需要在实验章明确取值依据并通过范围敏感性检验评估结论是否依赖特定参数。尤其是‘价值背叛—抵制—激进发帖’不能由该报告直接推出，若将其作为漂绿情境下的类型差异假设，仍需补充漂绿、品牌伪善与消费者表达行为研究【此处需要补充材料】。",
    ])

    replace_section(doc, "信息单元的模型定位", "情绪效价与唤醒强度", [
        "模型中的信息单元首先是可被 Agent 读取的自然语言消息，而不是预先编码的多维数值向量。当前代码使用统一消息结构保存信息来源（source）、文本内容（content）和消息类型（type）；社交消息另保存发送者标识（sender_id）。全局新闻、企业澄清和社交帖子均通过这一结构进入感知与反思环节。",
        "消息到达 Agent 后，感知层将其写入本轮观察集合。Reflect 层不会为每条消息分别生成效价、唤醒、可信度、证据强度和主题一致性分数，而是将消息文本、Persona、当前状态与相关记忆共同提供给 LLM，要求其形成综合语义判断。因此，这些概念在当前实现中是提示词层面的判断维度，不是独立存储、可逐项回归或直接统计的状态变量。",
    ])
    replace_section(doc, "情绪效价与唤醒强度", "可信度、证据强度与主题一致性", [
        "情绪效价与唤醒强度用于约束 LLM 对文本影响方向和反应程度的综合理解。负面新闻、消费者质疑或企业澄清可能诱发不同方向和幅度的反应，但当前代码不分别输出效价值与唤醒值，而是将二者综合反映在原始情绪性信任冲击 trust_change_affective 中。",
        "原始冲击由 Reflect 层生成，并被限制在 −2.0 至 1.5 的研究者设定区间。该非对称区间用于控制单次负面冲击与正向修复的最大幅度，不是 Forrester 报告或现有材料给出的经验估计。其合理性需要在实验章通过边界与敏感性检验验证【此处需要补充材料】。",
    ])
    replace_section(doc, "可信度、证据强度与主题一致性", "消息类型与信源差异", [
        "可信度、证据强度与主题一致性用于指导 LLM 判断企业声明是否直接回应争议、是否提供可核验依据，以及社交意见是否具有足够信息支持。当前实现只在 Prompt 中要求模型考虑这些维度，并未把三者写入独立字段或状态变量。正文因而不把它们表述为已经完成显式测量的数值机制。",
        "Reflect 层的结构化输出包括四项：是否感知企业伪善（hypocrisy_perceived）、原始情绪性信任冲击（trust_change_affective）、记忆重要性（importance）和自然语言判断理由（reasoning）。其中，原始冲击进入确定性信任更新公式，重要性进入记忆存储与检索，伪善判断和理由用于解释与日志追踪。若后续需要检验可信度或证据强度的独立效应，必须扩展输出模式和日志字段，而不能直接利用当前结果作事后推断。",
    ])
    replace_section(doc, "消息类型与信源差异", "UGC 二次生成与语义变形", [
        "当前实现区分全局新闻、企业澄清和社交帖子三类主要信息来源。Reflect 层在同一时间步最多组织第一条全局新闻、第一条企业澄清和两条社交帖子，并对文本长度进行截断后形成综合输入。该顺序是确定性工程规则，用于限制上下文规模，并不等同于基于可信度或重要性的动态排序。",
        "不同信源的解释差异通过 Prompt 实现：全局新闻作为外部事件信息，企业澄清需要结合其是否回应争议及是否提供证据进行判断，社交帖子作为其他消费者生成的二手意见进入语义环境。上述规则限定了 LLM 的判断视角，但不预设任一信源必然可信或必然产生固定方向的信任变化。",
    ])
    replace_section(doc, "UGC 二次生成与语义变形", "社交网络与传播机制", [
        "消费者在 Plan 与 Invoke 环节可能生成新的社交文本。该文本不是对原始消息的机械复制，而是由消费者类型、当前信任和本轮语义判断共同约束的再表达；生成后作为 social_review 消息沿有向边进入下游节点。",
        "当前代码能够保留再生成后的文本，但尚未记录 message_id、parent_id、传播跳数、原始证据保留率或文本语义距离。因此，模型可以观察不同 Agent 生成内容的文本差异，却不能据此严格识别‘证据衰减’‘主题漂移’或逐跳语义变形。上述指标如要成为论文中的实证机制，需要在消息模式、传播日志和指标模块中另行实现；在当前版本中仅作为可扩展方向。",
    ])

    replace_section(doc, "有向 BA 无标度网络构建", "节点角色与结构影响力", [
        "结构层采用有向化 BA 网络。系统先依据固定随机种子生成参数为 n 和 m=2 的无向 BA 底图，再计算底图节点度数，并将每条无向边定向为由较高度节点指向较低度节点；度数相同时按确定性的边遍历顺序定向。由此得到的有向图保留 BA 底图的异质连接结构，同时将边解释为消息从发送节点流向下游接收节点的传播通道。",
        "这一构造方式更准确地说是‘BA 底图的有向化’，而不是由有向优先连接过程直接生成的原生有向 BA 模型。该边界需要在第四章实现说明中明确。固定随机种子用于保证同一实验条件下网络结构可复现；不同网络种子则用于检验结论是否依赖单一拓扑实例。",
        "网络中的出度表示节点可直接触达的下游节点数量，入度表示节点可从多少上游节点接收信息。信息只沿有向边传播，因此方向相反的两个节点不存在自动双向互动。该设定适用于考察非对称广播路径，但不能直接代表具有互粉关系、社群重叠或算法推荐的完整平台网络。",
    ])
    replace_section(doc, "节点角色与结构影响力", "社交广播与级联传播", [
        "为避免概念混淆，本文统一使用‘Hub 节点’和‘普通节点’描述网络结构角色。Hub 节点由有向图中的出度确定，指能够直接触达较多下游节点的高出度节点；普通节点是相对低出度节点。Hub 是仿真网络内的拓扑属性，不等同于现实身份意义上的意见领袖，也不由 Persona 随机指定。",
        "消费者的社交活跃风格与结构角色相互独立。前者用于约束 Agent 是否倾向公开表达，后者决定帖子生成后能够触达多少下游节点。由此，同一 Agent 可能具有较强发帖倾向但处于普通节点，也可能位于 Hub 节点但本轮选择沉默。该分离使模型能够区分‘是否表达’与‘表达后能触达多少人’。",
        "渠道策略中的 Hub 投放据此选择出度最高的 top-K 节点，随机投放则在全部节点中按固定随机种子抽取 K 个节点。Hub 投放提高的是初始直接可达范围，而非预设说服效果；澄清是否改变信任仍取决于 Agent 的语义判断和确定性信任更新。",
    ])

    replace_by_prefix(doc, "语义层用于处理竞争性信息环境", "语义层用于将自然语言信息转化为可进入确定性信任更新的结构化判断。当前实现不为效价、唤醒、可信度、证据强度和主题一致性分别赋值，而是在 Prompt 中将其作为综合判断维度，由 Reflect 层输出企业伪善感知、原始情绪性信任冲击、记忆重要性和判断理由。其中，只有原始冲击进入信任更新公式，重要性用于记忆检索，其余输出用于解释和日志追踪。")
    replace_by_prefix(doc, "每一时间步中，系统首先注入", "每一时间步中，系统首先注入外部事件或企业澄清信息；Agent 随后读取信箱并形成观察集合。Reflect 层按固定工程规则组织第一条全局新闻、第一条企业澄清和至多两条社交帖子，并结合 Persona、当前状态与相关记忆生成结构化语义判断。Plan 层读取其中的原始情绪性信任冲击，依据消费者类型敏感度、既有信任、冲击锚点和恢复规则更新 trust_score，再由 LLM 在已确定的信任状态下判断购买、发帖及文本表达。最后，帖子沿有向边进入下游节点，在下一时间步被处理。")
    replace_by_prefix(doc, "第二，语义冲击是信任更新的直接输入", "第二，Reflect 层输出的原始情绪性信任冲击是确定性信任更新的直接语义输入。自然语言信息的来源、内容及其与消费者 Persona 和记忆的关系共同影响这一输出；效价、唤醒、可信度、证据强度和主题一致性在当前实现中属于 Prompt 判断维度，而非独立状态变量。该假设支撑模型使用 GABM 处理文本差异，同时保留 trust_score 的公式化更新。")
    replace_by_prefix(doc, "因此，消费者智能体在模型中同时承担三类功能", "因此，消费者智能体在模型中同时承担三类功能：其一，作为心理主体，承载消费者类型、信任和记忆等状态；其二，作为语义主体，依据 Persona、观察集合和相关记忆形成结构化判断；其三，作为网络主体，在有向结构中接收信息，并在选择发帖后向下游节点传播文本。")
    replace_by_prefix(doc, "本研究采用“生成式语义评估—确定性状态更新”", "本研究采用“生成式语义评估—确定性状态更新”的分层架构。Reflect 层负责组织当前观察、检索相关记忆并生成结构化语义判断；Plan 层负责读取原始情绪性信任冲击，并依据确定性规则更新最终信任值，再在更新后的状态下调用 LLM 形成购买与发帖判断。二者具有明确职责边界：Reflect 层不直接写入 trust_score，Plan 层才执行信任状态更新。该划分是程序职责与可解释性设计，不将其直接等同于未经验证的双过程心理结构。")
    replace_by_prefix(doc, "当观察集合中存在多条信息时", "当观察集合中存在多条信息时，Reflect 层按固定顺序组织输入：最多取第一条全局新闻、第一条企业澄清和两条社交帖子，并对文本长度进行截断。该规则用于控制单次 LLM 上下文，不构成来源可信度或信息重要性的动态排序。记忆模块另从历史记录中检索相关度较高的三条记忆作为上下文。")
    replace_by_prefix(doc, "Reflect 层最终输出原始情绪冲击量", "Reflect 层最终输出由企业伪善感知、原始情绪性信任冲击、记忆重要性和判断理由构成的 JSON 对象。原始冲击被限制在 −2.0 至 1.5，重要性被限制在 1 至 10；边界用于约束生成波动，属于待检验的建模参数。")
    replace_by_prefix(doc, "该范围设置体现负面偏差原则", "原始冲击的非对称范围意味着模型允许单次负面冲击的最大绝对幅度高于正向冲击。该设定用于表达‘损耗可能快于修复’的研究假设，而非现有材料已经估计出的心理参数；后续需要通过替代边界与对称区间实验检验结论稳健性【此处需要补充材料】。")
    replace_by_prefix(doc, "第三阶段是认知层（reflect）", "第三阶段是认知层（reflect），即结构化语义判断生成。Agent 根据固定规则组织观察文本，并结合 Persona、当前状态和相关记忆，输出企业伪善感知、原始情绪性信任冲击、记忆重要性和判断理由。当前实现不单独生成效价、唤醒、可信度、证据强度或主题一致性分数。")
    replace_by_prefix(doc, "语义层是心理层的输入转换机制", "语义层是心理层的输入转换机制。外部事件、企业澄清和社交帖子以自然语言文本进入 Reflect 层，由 LLM 在 Persona、当前状态和记忆约束下生成结构化判断。原始情绪性信任冲击进入确定性信任公式，重要性进入记忆机制，伪善感知与判断理由用于解释和追踪。")
    replace_by_prefix(doc, "心理层与语义层之间的耦合体现为", "心理层与语义层之间的耦合体现为：相同文本在不同消费者类型、既有信任与历史记忆条件下可能产生不同的原始情绪性信任冲击。语义层提供对文本的条件化判断，心理层提供 Persona 与状态边界，二者共同决定进入确定性信任更新的输入。")
    replace_by_prefix(doc, "语义层与结构层之间的耦合体现在", "语义层与结构层之间的耦合体现在：Agent 生成的社交文本沿有向边进入下游节点，并在下一时间步成为新的语义输入。当前模型能够保留再生成文本，但尚未实现传播谱系和语义距离指标，因此只讨论文本再表达与扩散，不把证据衰减或主题漂移写成已经测量的结果。")
    replace_by_prefix(doc, "再次，本章将自然语言信息转化为", "再次，本章将自然语言信息转化为可进入信任更新的结构化语义判断。当前消息对象保存来源、内容、类型和可选发送者，Reflect 层输出企业伪善感知、原始情绪性信任冲击、记忆重要性和判断理由。效价、唤醒、可信度、证据强度与主题一致性仅作为 Prompt 判断维度，不作为已独立量化的消息属性。")
    replace_by_prefix(doc, "将社交网络纳入模型具有明确的理论必要性", "将社交网络纳入模型具有明确的理论必要性。绿色消费扩散并非孤立个体决策的简单加总，而是在信息传播与社会互动中形成的动态过程。消费者不仅接收企业发布的信息，也会观察其他用户的评价和行为；高连接节点的示范与普通节点的持续互动均可能改变信息覆盖和讨论轨迹。由此，漂绿危机中的信任演化不能只被理解为“外部事件—个体反应”的线性过程，而应被置于社交传播关系中分析。")
    replace_by_prefix(doc, "社交广播的理论基础来自信息级联", "社交广播机制用于描述个体生成内容如何进入其他消费者的信息环境。消费者生成帖子后，系统沿发送节点的出向边将消息写入下游节点信箱；下游节点在下一时间步读取该消息并重新进行语义判断。跨 Tick 传播既形成可追踪的级联层次，也避免同一时间步内出现递归扩散。")
    replace_by_prefix(doc, "社交广播机制使信息传播具有明显的结构依赖性", "社交广播机制使信息传播具有明显的结构依赖性。若初始发帖者是 Hub 节点，其内容能够直接触达较多下游节点，并可能在下一时间步激活二次发帖；若初始发帖者是普通节点，其内容即使反应强烈，也可能只影响少数下游节点。因此，信息能否形成较大级联，同时取决于 Agent 是否发帖和发送节点的出度。")
    replace_by_prefix(doc, "社交网络机制在允许信息级联扩散的同时", "社交网络机制在允许信息级联扩散的同时，也需要限制单个 Agent 每轮处理的信息规模。当前代码在 Reflect 层最多组织第一条全局新闻、第一条企业澄清和两条社交帖子；在 Tick 末端，主实验路径另按到达顺序保留最新三条待处理消息。上述规则是固定容量与截断机制，不是注意力得分排序。第四章应如实说明，并在实验章检验容量变化是否改变结论。")
    replace_by_prefix(doc, "社会影响与口碑传播研究进一步说明", "渠道节点的结构位置会改变澄清的初始可达范围。企业澄清并不是发布后自动全网扩散，其后续覆盖取决于初始接收节点的出度、该节点是否生成帖子以及下游节点是否继续表达。Hub 节点的优势是较大的一阶触达范围，而不是天然具有更高可信度或更强说服力。")

    # 在策略边界中明确实验刺激与事件时间线的约束。
    boundary = find_heading(doc, "澄清注入规则与策略边界")
    cursor = boundary
    for _ in range(7):
        cursor = Paragraph(cursor._p.getnext(), cursor._parent)
    paragraph_after(cursor, "企业澄清文本中的审计比例、资金规模、认证主体和整改承诺等具体事实必须来自已授权且可核验的材料；缺少来源时只能使用不含虚构事实的中性实验文本【此处需要补充材料】。此外，主效应实验应围绕单一漂绿危机比较内容、渠道和时机，避免在修复窗口继续注入健康争议、金融风险等异质事件；多事件时间线应作为独立稳健性情景，而非主实验。")

    doc.save(OUTPUT)
    strip_comments(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
